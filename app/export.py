"""Exports a batch of invoices to an Excel workbook (Invoices, Line items, Issues)."""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Invoice

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
REVIEW_FILL = PatternFill("solid", fgColor="FDECEA")
MONEY = '"$"#,##0.00'


def _sheet(wb, title, headers, rows, money_cols=(), widths=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for col in money_cols:
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = MONEY
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _d(iso: str | None):
    try:
        return date.fromisoformat(iso) if iso else None
    except ValueError:
        return iso


def to_excel(invoices: list[Invoice]) -> bytes:
    wb = Workbook()
    rows = [[i.file_name, i.vendor, i.invoice_number, _d(i.invoice_date), _d(i.due_date),
             i.subtotal, i.tax, i.total, i.status, "; ".join(x.message for x in i.issues), i.extracted_by]
            for i in invoices]
    ws = _sheet(wb, "Invoices",
                ["File", "Vendor", "Invoice #", "Invoice date", "Due date", "Subtotal", "Tax", "Total",
                 "Status", "Issues", "Extracted by"],
                rows, money_cols=(6, 7, 8), widths=[32, 26, 14, 13, 13, 12, 10, 12, 13, 70, 13], first=True)
    for r in range(2, ws.max_row + 1):
        for c in (4, 5):
            ws.cell(r, c).number_format = "yyyy-mm-dd"
        if ws.cell(r, 9).value == "Review":
            for c in range(1, 12):
                ws.cell(r, c).fill = REVIEW_FILL
        ws.cell(r, 10).alignment = Alignment(wrap_text=True, vertical="top")

    last = ws.max_row
    total_row = last + 2
    ws.cell(total_row, 7, "Total").font = Font(bold=True)
    ws.cell(total_row, 8, f"=SUM(H2:H{max(last, 2)})").number_format = MONEY
    ws.cell(total_row, 8).font = Font(bold=True)

    _sheet(wb, "Line items", ["File", "Vendor", "Invoice #", "Description", "Qty", "Unit price", "Amount"],
           [[i.file_name, i.vendor, i.invoice_number, li.description, li.quantity, li.unit_price, li.amount]
            for i in invoices for li in i.line_items],
           money_cols=(6, 7), widths=[32, 26, 14, 40, 8, 12, 12])

    _sheet(wb, "Issues", ["File", "Vendor", "Invoice #", "Level", "Issue"],
           [[i.file_name, i.vendor, i.invoice_number, x.level, x.message] for i in invoices for x in i.issues],
           widths=[32, 26, 14, 10, 90])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
