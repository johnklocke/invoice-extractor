import io
import os
from datetime import date
from pathlib import Path
from types import SimpleNamespace

os.environ["ANTHROPIC_API_KEY"] = ""  # tests use the offline text rules

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.checks import check_batch  # noqa: E402
from app.export import to_excel  # noqa: E402
from app.extract import extract, parse_claude, parse_text, to_iso_date  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Invoice, LineItem  # noqa: E402

ROOT = Path(__file__).parent.parent
INVOICES = sorted((ROOT / "sample_data" / "invoices").glob("*.pdf"))
TODAY = date(2026, 10, 24)


def batch():
    return check_batch([extract(p.read_bytes(), p.name) for p in INVOICES], today=TODAY)


def test_dates_normalize():
    assert to_iso_date("Oct 3, 2026") == "2026-10-03"
    assert to_iso_date("10/07/2026") == "2026-10-07"
    assert to_iso_date("September 27, 2026") == "2026-09-27"
    assert to_iso_date("2026-10-12") == "2026-10-12"
    assert to_iso_date("soon") is None


def test_extracts_fields_from_samples():
    invs = {i.file_name: i for i in batch()}
    acme = invs["01_acme_office_supply.pdf"]
    assert (acme.vendor, acme.invoice_number, acme.invoice_date, acme.total) == \
        ("Acme Office Supply", "INV-1042", "2026-10-03", 471.60)
    assert len(acme.line_items) == 3 and acme.line_items[1].amount == 179.00

    nw = invs["02_northwind_plumbing.pdf"]  # different labels: "Invoice No.", "Payment Due", "Amount Due"
    assert (nw.invoice_number, nw.due_date, nw.tax, nw.total) == ("2231", "2026-11-06", 50.58, 773.08)
    assert nw.line_items[2].quantity == 2.5


def test_checks_catch_the_planted_problems():
    invs = {i.file_name: i for i in batch()}
    status = {name: inv.status for name, inv in invs.items()}
    assert status["01_acme_office_supply.pdf"] == "OK"
    assert status["02_northwind_plumbing.pdf"] == "OK"
    assert status["06_greenway_food_distributors.pdf"] == "OK"
    assert "add up to $1,022.00" in invs["03_brightline_electric.pdf"].issues[0].message
    assert "duplicate of 01_acme_office_supply.pdf" in invs["04_acme_office_supply_resent.pdf"].issues[0].message
    lakeside = [x.message for x in invs["05_lakeside_janitorial.pdf"].issues]
    assert any("total says $1,985.10" in m for m in lakeside)
    assert any(m.startswith("Past due") for m in lakeside)


def test_line_math_and_missing_fields():
    inv = Invoice(file_name="x.pdf", vendor="V", total=10.0,
                  line_items=[LineItem("Widget", 3, 2.00, 7.00)])
    check_batch([inv], today=TODAY)
    msgs = " ".join(x.message for x in inv.issues)
    assert "3 × $2.00 ≠ $7.00" in msgs and "Missing invoice number" in msgs and "Missing invoice date" in msgs


def test_scanned_pdf_and_images_need_ai_mode():
    assert "No text found" in parse_text("", "scan.pdf").issues[0].message
    img = ROOT / "sample_data" / "receipts" / "corner_hardware_receipt.png"
    inv = extract(img.read_bytes(), img.name)
    assert inv.status == "Review" and "AI mode" in inv.issues[0].message
    assert "Unsupported" in extract(b"hello", "notes.txt").issues[0].message


def test_claude_parser_maps_tool_output():
    fake_output = {"vendor": "Corner Hardware", "invoice_number": None, "invoice_date": "2026-10-18",
                   "due_date": None, "currency": "USD", "subtotal": 36.73, "tax": 2.57, "total": 39.30,
                   "line_items": [{"description": "Mouse traps 4pk", "quantity": 2, "unit_price": 7.99, "amount": 15.98}]}
    seen = {}

    class FakeClient:
        class messages:  # noqa: N801
            @staticmethod
            def create(**kwargs):
                seen.update(kwargs)
                return SimpleNamespace(content=[SimpleNamespace(type="tool_use", input=fake_output)])

    inv = parse_claude(b"\x89PNG...", "r.png", "image/png", client=FakeClient())
    assert seen["messages"][0]["content"][0]["type"] == "image"
    assert (inv.vendor, inv.total, inv.extracted_by) == ("Corner Hardware", 39.30, "claude")
    assert inv.line_items[0].amount == 15.98


def test_excel_export():
    wb = load_workbook(io.BytesIO(to_excel(batch())))
    assert wb.sheetnames == ["Invoices", "Line items", "Issues"]
    ws = wb["Invoices"]
    assert ws["B2"].value == "Acme Office Supply" and ws["I4"].value == "Review"
    assert wb["Line items"].max_row == 1 + 3 + 3 + 3 + 3 + 2 + 6


def test_web_and_api():
    client = TestClient(app)
    assert client.get("/").status_code == 200
    r = client.post("/samples")
    assert r.status_code == 200 and "need review" in r.text and "Download Excel" in r.text

    files = [("files", (p.name, p.read_bytes(), "application/pdf")) for p in INVOICES[:2]]
    data = client.post("/api/extract", files=files).json()
    assert [i["status"] for i in data["invoices"]] == ["OK", "OK"]
    xlsx = client.get(data["excel_url"])
    assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
    assert client.get("/batches/nope.xlsx").status_code == 404
